"""Filling an OGSERA template. `results.ogsera`, `ogsera_api`.

The template here is made up, shaped like OGSERA's: a title block of merged
cells above the header row, a learner's ID column, a name, the four Ogun
papers, a CA total that is a formula, an exam, a total, two traits and a
column OGSERA wants and Classnode does not fill ("Remark"); a data
validation on the exam column, a fill colour and a frozen pane, a second
sheet of instructions, and a footer row under the children.

Two schools in every test: the mapping is a row in each school's own schema.
"""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from academics.models import StudentDetails
from gradebook import services as gradebook
from gradebook.models import Assessment
from results import ogun
from results.models import OgseraMapping, Trait, TraitRating
from results.tests.fixtures import HOST, THEIR_HOST, ChainSetUp
from schools.tests.tenants import connected_to

BASE = "/api/results/ogsera/"

HEADINGS = [
    "S/N", "Learner ID", "Name of Learner", "1st Test (10)", "2nd Test (10)",
    "Assignment (10)", "CA Total (30)", "Exam (70)", "Total (100)",
    "Punctuality", "Neatness", "Remark",
]


def ogsera_template(rows, *, title="OGUN STATE MINISTRY OF EDUCATION, SCIENCE AND TECHNOLOGY"):
    """`rows`: (learner id, name) pairs. Returns the workbook's bytes."""
    book = Workbook()
    sheet = book.active
    sheet.title = "JSS1A Mathematics"
    sheet["A1"] = title
    sheet.merge_cells("A1:L1")
    sheet["A2"] = "OGSERA Continuous Assessment Template"
    sheet.merge_cells("A2:L2")
    sheet["A3"] = "School: St Mary's   Class: JSS 1A   Subject: Mathematics"
    sheet.append([])
    sheet.append(HEADINGS)  # row 5
    for cell in sheet[5]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFFF00")
    for n, (learner_id, name) in enumerate(rows, start=1):
        line = 5 + n
        sheet.append([n, learner_id, name, None, None, None, f"=SUM(D{line}:F{line})", None, None, None, None, None])
    # The table ends at its first empty row; a footer under it is no child.
    sheet.append([])
    sheet.append([None, None, "Class teacher's signature: ________"])
    exam = DataValidation(type="whole", operator="between", formula1="0", formula2="70")
    exam.add(f"H6:H{5 + max(len(rows), 1)}")
    sheet.add_data_validation(exam)
    sheet.freeze_panes = "D6"
    notes = book.create_sheet("Instructions")
    notes["A1"] = "Do not change the headings."
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def upload(raw, name="JSS1A_Mathematics.xlsx"):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(name, raw, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class OgseraSetUp(ChainSetUp):
    def setUp(self):
        super().setUp()
        self.ada, self.emeka = self.children["ada"], self.children["emeka"]
        with connected_to(self.stmarys):
            ogun.set_template_as(self.head.user, self.stmarys, "ogun")
            for child, learner_id in ((self.ada, "OG/0001"), (self.emeka, "OG/0002"), (self.children["bisi"], "OG/0003")):
                StudentDetails.objects.create(student_membership_id=child.pk, learner_id=learner_id)
            papers = {a.name: a for a in Assessment.objects.filter(term=self.term, subject=self.maths)}
            for name, value in (("1st Test", 8), ("2nd Test", 9), ("Assignment", 7), ("Exam", 61)):
                gradebook.set_score(papers[name], self.ada, value, by=self.teacher.user)
            self.punctuality = Trait.objects.get(name="Punctuality")
            TraitRating.objects.create(term=self.term, trait=self.punctuality, student_membership_id=self.ada.pk, score=4)
        with connected_to(self.grace):
            ogun.set_template_as(self.their_head.user, self.grace, "ogun")

    def post(self, route, raw, who=None, host=HOST, **data):
        self.client.force_login((who or self.head).user)
        return self.client.post(f"{BASE}{route}", {"file": upload(raw), **data}, HTTP_HOST=host)

    def map_columns(self, who=None, host=HOST):
        self.client.force_login((who or self.head).user)
        with connected_to(self.stmarys if host == HOST else self.grace):
            punctuality = Trait.objects.get(name="Punctuality").pk
        columns = {
            "Learner ID": "learner_id", "Name of Learner": "name", "1st Test (10)": "test1",
            "2nd Test (10)": "test2", "Assignment (10)": "assignment", "CA Total (30)": "ca_total",
            "Exam (70)": "exam", "Total (100)": "total", "Punctuality": f"trait:{punctuality}",
            "S/N": "", "Remark": "",
        }
        return self.client.put(f"{BASE}mapping/", data={"columns": columns}, content_type="application/json", HTTP_HOST=host)

    def good_file(self):
        return ogsera_template([("OG/0001", "Ada Obi"), ("og/0002", "Emeka Nwosu"), ("OG/0003", "Bisi Ade")])

    def form(self, **over):
        data = {"class_group_id": self.jss1a.pk, "subject_id": self.maths.pk}
        data.update(over)
        return data


class TheHeadingsTests(OgseraSetUp):
    def test_the_header_row_is_found_under_the_title_block(self):
        body = self.post("headings/", self.good_file()).json()

        self.assertEqual((body["sheet_name"], body["header_row"]), ("JSS1A Mathematics", 5))
        self.assertEqual([h["heading"] for h in body["headings"]], HEADINGS)
        self.assertTrue(all(h["field"] == "" for h in body["headings"]), "nothing is mapped yet")

    def test_a_saved_mapping_comes_back_on_the_headings(self):
        self.assertEqual(self.map_columns().status_code, 200)
        body = self.post("headings/", self.good_file()).json()
        fields = {h["heading"]: h["field"] for h in body["headings"]}
        self.assertEqual((fields["Learner ID"], fields["Exam (70)"], fields["Remark"]), ("learner_id", "exam", ""))

    def test_a_mapping_needs_exactly_one_learner_id_column(self):
        self.client.force_login(self.head.user)
        response = self.client.put(
            f"{BASE}mapping/", data={"columns": {"Name of Learner": "name"}}, content_type="application/json", HTTP_HOST=HOST
        )
        self.assertEqual(response.status_code, 422)
        self.assertIn("learner's ID", response.json()["detail"])


class TheCheckTests(OgseraSetUp):
    def setUp(self):
        super().setUp()
        self.map_columns()

    def test_the_control_a_good_file_matches_every_row_case_blind(self):
        body = self.post("check/", self.good_file(), **self.form()).json()

        self.assertEqual((body["matched"], body["unmatched"], body["missing_id"]), (3, [], []))
        self.assertTrue(body["may_download"])
        self.assertEqual(body["not_in_file"], [])
        self.assertEqual(body["no_learner_id"], ["Tunde Cole"])
        self.assertIn("Emeka Nwosu: no Exam", body["missing_values"])

    def test_an_unknown_id_and_a_row_with_no_id_block_the_download(self):
        raw = ogsera_template([("OG/0001", "Ada Obi"), ("OG/9999", "Stranger"), ("", "No ID Child")])
        body = self.post("check/", raw, **self.form()).json()

        self.assertEqual(body["unmatched"], [{"row": 7, "learner_id": "OG/9999"}])
        self.assertEqual(body["missing_id"], [{"row": 8, "name": "No ID Child"}])
        self.assertFalse(body["may_download"])
        self.assertEqual(body["not_in_file"], ["Bisi Ade", "Emeka Nwosu"])
        self.assertEqual(self.post("fill/", raw, **self.form()).status_code, 422)

    def test_a_child_of_another_class_is_unmatched_here(self):
        with connected_to(self.stmarys):
            StudentDetails.objects.create(student_membership_id=self.bimpe.pk, learner_id="OG/0005")
        body = self.post("check/", ogsera_template([("OG/0005", "Bimpe Ojo")]), **self.form()).json()
        self.assertEqual(body["unmatched"], [{"row": 6, "learner_id": "OG/0005"}])

    def test_a_marks_file_needs_its_subject(self):
        response = self.post("check/", self.good_file(), class_group_id=self.jss1a.pk)
        self.assertEqual(response.status_code, 422)
        self.assertIn("Choose the subject", response.json()["detail"])


class TheFillTests(OgseraSetUp):
    def setUp(self):
        super().setUp()
        self.map_columns()

    def filled(self, raw=None, **over):
        response = self.post("fill/", raw or self.good_file(), **self.form(**over))
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertIn("filled.xlsx", response["Content-Disposition"])
        return load_workbook(io.BytesIO(response.content))

    def test_the_control_mapped_cells_of_matched_rows_are_filled(self):
        sheet = self.filled()["JSS1A Mathematics"]

        self.assertEqual([sheet.cell(row=6, column=c).value for c in (4, 5, 6, 8, 9, 10)], [8, 9, 7, 61, 85, 4])
        self.assertEqual(sheet["C6"].value, "Ada Obi")

    def test_the_files_formatting_formula_validation_and_sheets_are_kept(self):
        book = self.filled()
        sheet = book["JSS1A Mathematics"]

        self.assertEqual(sheet["G6"].value, "=SUM(D6:F6)", "a formula cell is not overwritten")
        self.assertIn("A1:L1", [str(r) for r in sheet.merged_cells.ranges])
        self.assertTrue(sheet["A5"].font.bold)
        self.assertEqual(sheet["A5"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(sheet.freeze_panes, "D6")
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        self.assertEqual(book.sheetnames, ["JSS1A Mathematics", "Instructions"])

    def test_unmapped_columns_are_left_alone(self):
        sheet = self.filled()["JSS1A Mathematics"]
        self.assertIsNone(sheet["L6"].value, "Remark is not mapped")
        self.assertEqual(sheet["A6"].value, 1, "S/N is not mapped")
        self.assertIsNone(sheet["K6"].value, "Neatness is not mapped")

    def test_a_cell_with_a_value_is_kept_unless_replace_is_ticked(self):
        book = load_workbook(io.BytesIO(self.good_file()))
        book["JSS1A Mathematics"]["H6"] = 50
        out = io.BytesIO()
        book.save(out)

        kept = self.filled(out.getvalue())["JSS1A Mathematics"]
        replaced = self.filled(out.getvalue(), replace="1")["JSS1A Mathematics"]

        self.assertEqual(kept["H6"].value, 50)
        self.assertEqual(replaced["H6"].value, 61)

    def test_a_missing_mark_leaves_its_cell_empty(self):
        sheet = self.filled()["JSS1A Mathematics"]
        self.assertIsNone(sheet["H7"].value, "Emeka has no exam mark")

    def test_a_paper_not_out_of_the_sheets_maximum_is_not_filled(self):
        """A school that chose Ogun after marking keeps an "Exam" out of 100 (presets never
        overwrite a marked paper). Its 61 is not 61 of 70, so it is not written under "Exam (70)"."""
        with connected_to(self.stmarys):
            Assessment.objects.filter(term=self.term, subject=self.maths, name="Exam").update(max_score=100)

        body = self.post("check/", self.good_file(), **self.form()).json()
        sheet = self.filled()["JSS1A Mathematics"]

        self.assertIn("Ada Obi: no Exam", body["missing_values"])
        self.assertIsNone(sheet["H6"].value, "Ada's 61 out of 100 is not an Exam out of 70")
        self.assertIsNone(sheet["I6"].value, "and so there is no total built on it")
        self.assertEqual([sheet.cell(row=6, column=c).value for c in (4, 5, 6)], [8, 9, 7], "the tests are the sheet's own")

    def test_another_schools_exam_paper_is_judged_by_that_schools_own_maximum(self):
        """Grace marks its Exam out of 100 and St Mary's out of 70: only St Mary's mark is filled,
        and changing Grace's paper does not move St Mary's."""
        with connected_to(self.grace):
            Assessment.objects.filter(subject__name="Mathematics", name="Exam").update(max_score=100)

        sheet = self.filled()["JSS1A Mathematics"]

        self.assertEqual(sheet["H6"].value, 61, "St Mary's Exam is still out of 70")
        with connected_to(self.stmarys):
            self.assertEqual(
                Assessment.objects.get(term=self.term, subject=self.maths, name="Exam").max_score, 70
            )


class WhoAndWhereTests(OgseraSetUp):
    def test_a_teacher_and_a_bursar_are_refused(self):
        for who in (self.teacher, self.bursar):
            with self.subTest(who=who.user.username):
                self.assertEqual(self.post("headings/", self.good_file(), who=who).status_code, 403)

    def test_a_school_not_on_the_ogun_card_is_told_so(self):
        with connected_to(self.stmarys):
            ogun.set_template_as(self.head.user, self.stmarys, "standard")
        response = self.post("headings/", self.good_file())
        self.assertEqual(response.status_code, 409)
        self.assertIn("Ogun State template", response.json()["detail"])

    def test_one_schools_mapping_is_never_the_others(self):
        self.map_columns()
        self.client.force_login(self.their_head.user)
        theirs = self.client.get(BASE, HTTP_HOST=THEIR_HOST).json()
        ours = self.client.get(BASE, HTTP_HOST=HOST)

        self.assertEqual(theirs["mapping"], [])
        self.assertEqual(ours.status_code, 403, "the other school's principal is refused here")
        with connected_to(self.grace):
            self.assertFalse(OgseraMapping.objects.exists())
        with connected_to(self.stmarys):
            self.assertEqual(OgseraMapping.objects.get().columns["learner id"], "learner_id")

    def test_the_other_schools_children_are_not_matched_from_their_host(self):
        self.map_columns()
        self.map_columns(who=self.their_head, host=THEIR_HOST)
        with connected_to(self.grace):
            StudentDetails.objects.create(student_membership_id=self.grace_child.pk, learner_id="OG/7777")
        body = self.post(
            "check/", self.good_file(), who=self.their_head, host=THEIR_HOST,
            class_group_id=self.grace_group.pk, subject_id=self.grace_maths_id(),
        ).json()
        self.assertEqual(body["matched"], 0)
        self.assertEqual(len(body["unmatched"]), 3)

    def grace_maths_id(self):
        from gradebook.models import Subject

        with connected_to(self.grace):
            return Subject.objects.get(name="Mathematics").pk
