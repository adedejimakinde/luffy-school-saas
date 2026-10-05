"""The roll import from a file: an Excel workbook or a CSV, checked, then admitted.

`test_bulk_enrolment.py` holds every rule of `accounts/bulk.py` against CSV text
posted as JSON. What is held here is what the upload adds on top of it, with
**two schools in every test** for that file's reason: `User` and `Membership`
are shared tables, so what keeps one school's import out of another's roll is
an authority check and a scoped lookup, not the tenant schema.

- a workbook is read into the same rows a CSV is, with the office's own row
  numbers, and a formatted-but-empty row is not a child;
- the preview is `check()`'s verdict on every row and **writes nothing**;
- the file import is all or nothing, as the CSV one is, and a duplicate
  admission number refuses the whole file by row;
- the template lists this school's classes, and only this school's;
- a child's learner's ID, sex and date of birth come in with them, checked by
  row, and a learner's ID is unique at a school as an admission number is.
"""

import datetime

import io

from openpyxl import Workbook, load_workbook

from academics.models import ClassGroup, ClassPlacement, StudentDetails
from accounts import bulk
from accounts.models import Membership, Role, User
from accounts.tests.test_enrolment_api import EnrolmentSetUp
from results.tests.fixtures import HOST, THEIR_HOST
from schools.tests.tenants import connected_to

BASE = "/api/enrolment/roll/import/"
DOOR = f"{BASE}door/"
CHECK = f"{BASE}check/"
FILE = f"{BASE}file/"
TEMPLATE = f"{BASE}template/"

HEADINGS = [
    "Full name",
    "Class group",
    "Reference",
    "Username",
    "Guardian name",
    "Guardian contact",
    "Learner ID",
    "Sex",
    "Date of birth",
]


def workbook(*rows, headings=HEADINGS, gap_rows=(), sheet_first=None):
    """A real .xlsx, as bytes. `gap_rows` are row numbers left formatted and empty."""
    book = Workbook()
    sheet = book.active
    if sheet_first is not None:
        sheet.title = "Notes"
        sheet.append([sheet_first])
        sheet = book.create_sheet("Students")
    sheet.append(headings)
    line = 2
    for row in rows:
        while line in gap_rows:
            sheet.cell(row=line, column=1).number_format = "@"
            line += 1
        for column, value in enumerate(row, start=1):
            sheet.cell(row=line, column=column, value=value)
        line += 1
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def upload(name, raw, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(name, raw, content_type=content_type)


class RollImportSetUp(EnrolmentSetUp):
    def post(self, route, member, raw, host=HOST, name="roll.xlsx"):
        self.client.force_login(member.user)
        return self.client.post(route, {"file": upload(name, raw)}, HTTP_HOST=host)

    def students(self, school):
        return Membership.objects.filter(school=school, role=Role.STUDENT).count()

    def counts(self):
        return (User.objects.count(), self.students(self.stmarys), self.students(self.grace))


class ReadingAWorkbookTests(RollImportSetUp):
    """`bulk.read_upload()`, on its own."""

    def test_a_workbook_reads_into_the_rows_a_csv_does(self):
        raw = workbook(["Ada Obi", "JSS 1A", "0100", "", "", ""])

        self.assertEqual(
            bulk._read(bulk.read_upload(raw)),
            bulk._read(
                "full_name,class_group,reference,username,guardian_name,guardian_contact,"
                "learner_id,sex,date_of_birth\nAda Obi,JSS 1A,0100,,,,,,\n"
            ),
        )

    def test_row_numbers_are_the_spreadsheets_and_empty_rows_are_nobody(self):
        raw = workbook(["Ada Obi", "JSS 1A"], ["Bola Ade", "JSS 1B"], gap_rows=(3, 4))

        self.assertEqual([line for line, _ in bulk._read(bulk.read_upload(raw))], [2, 5])

    def test_numbers_typed_into_general_cells_come_back_as_typed(self):
        """An admission number or a phone typed as a number is not 100.0."""
        raw = workbook(["Ada Obi", "JSS 1A", 100.0, None, "Mrs Obi", 8031234567])

        _, row = bulk._read(bulk.read_upload(raw))[0]
        self.assertEqual((row["reference"], row["guardian_contact"]), ("100", "8031234567"))

    def test_the_students_sheet_is_found_behind_a_sheet_of_notes(self):
        raw = workbook(["Ada Obi", "JSS 1A"], sheet_first="Fill in the next sheet")

        self.assertEqual(bulk._read(bulk.read_upload(raw))[0][1]["full_name"], "Ada Obi")

    def test_an_old_xls_and_a_non_spreadsheet_are_refused_with_a_sentence(self):
        for raw, said in (
            (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "older Excel file"),
            (b"PK\x03\x04not really a zip", "not a spreadsheet this can read"),
            (b"\x89PNG\r\n\x1a\n\x00\x00", "not a spreadsheet this can read"),
        ):
            with self.subTest(said=said), self.assertRaisesMessage(bulk.BulkError, said):
                bulk._read(bulk.read_upload(raw))

    def test_a_file_over_the_limit_is_refused_unopened(self):
        with self.assertRaisesMessage(bulk.BulkError, "over 2 MB"):
            bulk.read_upload(b"x" * (bulk.MAX_FILE_BYTES + 1))

    def test_a_workbook_that_unpacks_huge_is_refused_before_it_is_inflated(self):
        import zipfile

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"\x00" * (bulk.MAX_UNPACKED_BYTES + 1))
        with self.assertRaisesMessage(bulk.BulkError, "far larger than a roll needs"):
            bulk.read_upload(out.getvalue())

    def test_the_template_sent_straight_back_is_not_an_empty_success(self):
        with self.assertRaisesMessage(bulk.BulkError, "no children under them"):
            bulk._read(bulk.read_upload(bulk.template_workbook(["JSS 1A"])))


class ThePreviewTests(RollImportSetUp):
    def test_the_control_a_good_workbook_is_ready_on_every_row(self):
        body = self.post(CHECK, self.admin, workbook(
            ["Ada Obi", "JSS 1A", "0100"],
            ["Bola Ade", "JSS 1B", "0101", "", "Mrs Ade", "0803 123 4567"],
        )).json()

        self.assertTrue(body["admissible"])
        self.assertEqual(body["problem_rows"], 0)
        self.assertEqual([r["line"] for r in body["rows"]], [2, 3])
        self.assertEqual(body["rows"][1]["guardian_contact"], "0803 123 4567")
        self.assertEqual([r["problems"] for r in body["rows"]], [[], []])

    def test_it_judges_every_row_and_writes_nothing(self):
        before = self.counts()

        body = self.post(CHECK, self.admin, workbook(
            ["Ada Obi", "JSS 1A", "0100"],
            ["", "JSS 9Z", "0100"],
            ["Chi Eze", "JSS 1A", "0102", "", "Mr Eze", "not a number"],
        )).json()

        self.assertFalse(body["admissible"])
        self.assertEqual(body["problem_rows"], 2)
        by_line = {r["line"]: [(p["column"], p["detail"]) for p in r["problems"]] for r in body["rows"]}
        self.assertEqual(by_line[2], [])
        self.assertEqual(
            [column for column, _ in by_line[3]], ["full_name", "reference", "class_group"]
        )
        self.assertEqual([column for column, _ in by_line[4]], ["guardian_contact"])
        self.assertEqual(self.counts(), before)

    def test_even_a_clean_file_is_not_admitted_by_the_preview(self):
        before = self.counts()

        self.assertTrue(self.post(CHECK, self.admin, workbook(["Ada Obi", "JSS 1A", "0100"])).json()["admissible"])

        self.assertEqual(self.counts(), before)

    def test_an_admission_number_already_at_this_school_is_named_by_row(self):
        Membership.objects.filter(pk=self.children["ada"].pk).update(reference="0100")

        body = self.post(CHECK, self.admin, workbook(["Chike Obi", "JSS 1A", "0100"])).json()

        self.assertFalse(body["admissible"])
        self.assertIn("already an admission number at this school", body["rows"][0]["problems"][0]["detail"])

    def test_the_same_admission_number_at_the_other_school_is_no_obstacle(self):
        grace_child = Membership.objects.filter(school=self.grace, role=Role.STUDENT).first()
        Membership.objects.filter(pk=grace_child.pk).update(reference="0100")

        self.assertTrue(self.post(CHECK, self.admin, workbook(["Chike Obi", "JSS 1A", "0100"])).json()["admissible"])

    def test_a_class_only_the_other_school_teaches_is_not_a_class_here(self):
        with connected_to(self.grace):
            ClassGroup.objects.create(name="Grace Ruby", level=1)

        body = self.post(CHECK, self.admin, workbook(["Chike Obi", "Grace Ruby"])).json()
        theirs = self.post(CHECK, self.their_admin, workbook(["Chike Obi", "Grace Ruby"]), host=THEIR_HOST).json()

        self.assertFalse(body["admissible"])
        self.assertIn("not a class this school teaches", body["rows"][0]["problems"][0]["detail"])
        self.assertTrue(theirs["admissible"], "the control: Grace's own admin may use Grace's class")

    def test_a_file_that_is_not_a_spreadsheet_is_a_422_with_a_sentence(self):
        response = self.post(CHECK, self.admin, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, name="roll.xls")

        self.assertEqual(response.status_code, 422)
        self.assertIn("Excel Workbook (.xlsx)", response.json()["detail"])

    def test_no_file_at_all_is_a_422(self):
        self.client.force_login(self.admin.user)

        self.assertEqual(self.client.post(CHECK, {}, HTTP_HOST=HOST).status_code, 422)


class WhoMayImportTests(RollImportSetUp):
    """Every route asks before it reads, at the school whose host it is on."""

    def test_a_principal_a_teacher_and_a_bursar_are_refused_everywhere(self):
        raw = workbook(["Ada Obi", "JSS 1A"])
        for member in (self.head, self.teacher, self.bursar):
            for route in (CHECK, FILE):
                with self.subTest(member=member.user.username, route=route):
                    response = self.post(route, member, raw)
                    self.assertEqual(response.status_code, 403)
                    self.assertNotIn("rows", response.json())
            for route in (DOOR, TEMPLATE):
                with self.subTest(member=member.user.username, route=route):
                    self.client.force_login(member.user)
                    self.assertEqual(self.client.get(route, HTTP_HOST=HOST).status_code, 403)

    def test_an_administrator_at_one_school_is_refused_at_the_other(self):
        before = self.counts()
        raw = workbook(["Ada Obi", "JSS 1A"])

        for route in (CHECK, FILE):
            with self.subTest(route=route):
                self.assertEqual(self.post(route, self.admin, raw, host=THEIR_HOST).status_code, 403)
        self.client.force_login(self.admin.user)
        self.assertEqual(self.client.get(TEMPLATE, HTTP_HOST=THEIR_HOST).status_code, 403)
        self.assertEqual(self.client.get(DOOR, HTTP_HOST=THEIR_HOST).status_code, 403)
        self.assertEqual(self.counts(), before)

    def test_the_door_says_the_term_and_this_schools_classes(self):
        with connected_to(self.grace):
            ClassGroup.objects.create(name="Grace Ruby", level=1)
        self.client.force_login(self.admin.user)

        body = self.client.get(DOOR, HTTP_HOST=HOST).json()

        self.assertEqual(body["classes"], ["JSS 1A", "JSS 1B"])
        self.assertTrue(body["term"])


class TheFileImportTests(RollImportSetUp):
    def test_the_control_a_good_workbook_admits_everybody_and_places_them(self):
        body = self.post(FILE, self.admin, workbook(
            ["Chike Obi", "JSS 1A", "0100", "STM/2026/0100"],
            ["Ngozi Abah", "JSS 1B", "0101"],
        )).json()

        self.assertEqual((body["admitted"], body["problems"]), (2, []))
        self.assertEqual(body["generated"], {"3": "ST-MARYS/0101"})
        chike = Membership.objects.get(user__username="STM/2026/0100")
        self.assertEqual(chike.school, self.stmarys)
        self.assertEqual(chike.reference, "0100")
        with connected_to(self.stmarys):
            self.assertTrue(
                ClassPlacement.objects.filter(student_membership_id=chike.pk, class_group_id=self.jss1a_id).exists()
            )

    def test_a_duplicate_admission_number_refuses_the_whole_file_by_row(self):
        """The rule kept from the CSV import: not skipped, refused, by line."""
        before = self.counts()

        body = self.post(FILE, self.admin, workbook(
            ["Chike Obi", "JSS 1A", "0100"],
            ["Ngozi Abah", "JSS 1B", "0101"],
            ["Tobi Ade", "JSS 1B", "0100"],
        )).json()

        self.assertEqual(body["admitted"], 0)
        self.assertEqual([(p["line"], p["column"]) for p in body["problems"]], [(4, "reference")])
        self.assertEqual(self.counts(), before, "the two good rows were written")

    def test_an_import_at_one_school_leaves_the_other_untouched(self):
        grace_before = self.students(self.grace)

        self.post(FILE, self.admin, workbook(["Chike Obi", "JSS 1A", "0100"]))

        self.assertEqual(self.students(self.grace), grace_before)
        self.assertEqual(Membership.objects.get(user__full_name="Chike Obi").school, self.stmarys)

    def test_a_csv_file_goes_through_the_same_door(self):
        raw = b"full_name,class_group,reference\r\nChike Obi,JSS 1A,0100\r\n"

        body = self.post(FILE, self.admin, raw, name="roll.csv").json()

        self.assertEqual(body["admitted"], 1)

    def test_a_csv_saved_by_excel_in_the_windows_code_page_is_read(self):
        raw = "full_name,class_group\r\nAdébáyọ̀ Ọlá,JSS 1A\r\n".encode("utf-8")
        latin = "full_name,class_group\r\nRenée Adé,JSS 1A\r\n".encode("cp1252")

        self.assertEqual(self.post(FILE, self.admin, raw, name="a.csv").json()["admitted"], 1)
        self.assertEqual(self.post(FILE, self.admin, latin, name="b.csv").json()["admitted"], 1)
        self.assertTrue(User.objects.filter(full_name="Renée Adé").exists())


class TheTemplateTests(RollImportSetUp):
    def template(self, member, host=HOST):
        self.client.force_login(member.user)
        response = self.client.get(TEMPLATE, HTTP_HOST=host)
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        return load_workbook(io.BytesIO(response.content))

    def listed(self, book):
        return [row[0] for row in book["Classes"].iter_rows(values_only=True)]

    def test_each_school_gets_its_own_classes_and_only_those(self):
        with connected_to(self.grace):
            ClassGroup.objects.create(name="Grace Ruby", level=1)
        with connected_to(self.stmarys):
            ClassGroup.objects.create(name="JSS 4A", level=4, is_active=False)

        ours = self.listed(self.template(self.admin))
        theirs = self.listed(self.template(self.their_admin, host=THEIR_HOST))

        self.assertEqual(ours, ["JSS 1A", "JSS 1B"])
        self.assertIn("Grace Ruby", theirs)
        self.assertNotIn("Grace Ruby", ours)
        self.assertNotIn("JSS 1B", theirs)

    def test_it_has_the_headings_the_import_reads_and_no_example_child(self):
        book = self.template(self.admin)
        rows = list(book["Students"].iter_rows(values_only=True))

        self.assertEqual(list(rows[0]), HEADINGS)
        self.assertTrue(all(all(c is None for c in r) for r in rows[1:]))
        self.assertEqual(book["Students"]["C2"].number_format, "@", "admission numbers keep their zeros")
        self.assertEqual(book["Students"]["G2"].number_format, "@", "learner IDs keep their zeros")
        self.assertEqual(book["Students"]["I2"].number_format, "DD/MM/YYYY")
        self.assertEqual(book["Classes"].sheet_state, "hidden")
        self.assertIn("Classes!$A$1:$A$2", book["Students"].data_validations.dataValidation[0].formula1)

    def test_filled_in_it_is_a_file_the_import_takes(self):
        book = self.template(self.admin)
        book["Students"].append(["Chike Obi", "JSS 1A", "0100"])
        out = io.BytesIO()
        book.save(out)

        self.assertTrue(self.post(CHECK, self.admin, out.getvalue()).json()["admissible"])


class TheChildsDetailsTests(RollImportSetUp):
    """Learner's ID, sex and date of birth: three optional columns."""

    def test_the_control_they_are_admitted_with_the_child(self):
        body = self.post(FILE, self.admin, workbook(
            ["Chike Obi", "JSS 1A", "0100", "", "", "", "OG/ABS/0100", "M", datetime.datetime(2013, 5, 2)],
            ["Ngozi Abah", "JSS 1B", "0101", "", "", "", "", "Female", "02/03/2014"],
            ["Tobi Ade", "JSS 1B", "0102"],
        )).json()

        self.assertEqual((body["admitted"], body["problems"]), (3, []))
        with connected_to(self.stmarys):
            rows = {
                Membership.objects.get(pk=row.student_membership_id).user.full_name: row
                for row in StudentDetails.objects.all()
            }
        self.assertEqual(sorted(rows), ["Chike Obi", "Ngozi Abah"], "a row with none of it makes no details row")
        chike, ngozi = rows["Chike Obi"], rows["Ngozi Abah"]
        self.assertEqual((chike.learner_id, chike.sex, chike.date_of_birth), ("OG/ABS/0100", "male", datetime.date(2013, 5, 2)))
        self.assertEqual((ngozi.learner_id, ngozi.sex, ngozi.date_of_birth), ("", "female", datetime.date(2014, 3, 2)))

    def test_each_bad_value_is_named_by_row_and_column_and_nothing_is_written(self):
        before = self.counts()
        body = self.post(FILE, self.admin, workbook(
            ["Chike Obi", "JSS 1A", "", "", "", "", "OG/1", "boy", "31/02/2013"],
            ["Ngozi Abah", "JSS 1B", "", "", "", "", "og/1", "unknown", "2099-01-01"],
        )).json()

        self.assertEqual(body["admitted"], 0)
        self.assertEqual(
            [(p["line"], p["column"]) for p in body["problems"]],
            [(2, "date_of_birth"), (3, "learner_id"), (3, "sex"), (3, "date_of_birth")],
        )
        self.assertEqual(self.counts(), before)

    def test_a_learner_id_already_at_this_school_is_refused_and_at_the_other_is_not(self):
        with connected_to(self.stmarys):
            StudentDetails.objects.create(student_membership_id=self.children["ada"].pk, learner_id="OG/7")

        ours = self.post(CHECK, self.admin, workbook(["Chike Obi", "JSS 1A", "", "", "", "", "og/7"])).json()
        theirs = self.post(
            CHECK, self.their_admin, workbook(["Chike Obi", "JSS 1A", "", "", "", "", "OG/7"]), host=THEIR_HOST
        ).json()

        self.assertEqual([p["column"] for r in ours["rows"] for p in r["problems"]], ["learner_id"])
        self.assertIn("already a learner's ID at this school", ours["rows"][0]["problems"][0]["detail"])
        self.assertTrue(theirs["admissible"])

    def test_the_preview_shows_what_was_read(self):
        body = self.post(CHECK, self.admin, workbook(
            ["Chike Obi", "JSS 1A", "", "", "", "", "OG/1", "F", datetime.datetime(2013, 5, 2)],
        )).json()

        row = body["rows"][0]
        self.assertEqual((row["learner_id"], row["sex"], row["date_of_birth"]), ("OG/1", "F", "2013-05-02"))

    def test_the_template_offers_female_and_male(self):
        self.client.force_login(self.admin.user)
        book = load_workbook(io.BytesIO(self.client.get(TEMPLATE, HTTP_HOST=HOST).content))
        formulas = [v.formula1 for v in book["Students"].data_validations.dataValidation]
        self.assertIn('"Female,Male"', formulas)


class TheEmailGapTests(RollImportSetUp):
    """After an import: who will get no alert or receipt, because those go by email only.

    Two schools, because the gap is judged by **what this school may know**: a guardian who
    holds an email at another school is no email here until they have answered this one.
    """

    MUM = "0803 123 0000"

    def import_(self, *rows):
        return self.post(FILE, self.admin, workbook(*rows)).json()

    def gaps(self, body):
        return [(g["line"], g["full_name"], g["why"]) for g in body["no_email"]]

    def live_guardian_with_an_email(self, school, child, phone, email):
        """A guardian whose link at `school` is live and who holds `email` there."""
        from django.utils import timezone

        from accounts import guardian_contacts
        from accounts.models import ContactChannel, GuardianContact
        from accounts.services import activate_guardian_links

        admin = self.admin if school == self.stmarys else self.their_admin
        link = guardian_contacts.link_by_contact_as(admin, child, "Mrs Live", phone)
        account = link.guardian.guardian_account
        activate_guardian_links(link.guardian, school)
        GuardianContact.objects.create(
            guardian=account, channel_type=ContactChannel.EMAIL, value=email,
            created_by=admin.user, verified_at=timezone.now(),
        )
        return link.guardian

    def test_the_control_a_phone_or_nobody_is_listed_and_an_email_is_not(self):
        body = self.import_(
            ["Chike Obi", "JSS 1A", "0100", "", "Mrs Obi", "mum@example.test"],
            ["Ngozi Abah", "JSS 1B", "0101", "", "Mr Abah", self.MUM],
            ["Tobi Ade", "JSS 1A", "0102"],
        )

        self.assertEqual(body["admitted"], 3)
        self.assertEqual(
            self.gaps(body),
            [(3, "Ngozi Abah", "A phone number only."), (4, "Tobi Ade", "No guardian was given.")],
        )
        ngozi = body["no_email"][0]
        self.assertEqual(
            (ngozi["class_group"], ngozi["reference"], ngozi["guardian_name"], ngozi["guardian_contact"]),
            ("JSS 1B", "0101", "Mr Abah", "+2348031230000"),
        )

    def test_a_file_where_everyone_has_an_email_lists_nobody(self):
        body = self.import_(["Chike Obi", "JSS 1A", "0100", "", "Mrs Obi", "mum@example.test"])
        self.assertEqual(body["no_email"], [])

    def test_a_refused_file_lists_nobody_and_writes_nothing(self):
        before = self.counts()
        body = self.import_(
            ["Chike Obi", "JSS 1A", "0100", "", "Mrs Obi", self.MUM],
            ["Tobi Ade", "JSS 9Z", "0101"],
        )
        self.assertEqual((body["admitted"], body["no_email"]), (0, []))
        self.assertEqual(self.counts(), before)

    def test_a_guardian_already_live_here_with_an_email_is_not_a_gap(self):
        first = self.import_(["Chike Obi", "JSS 1A", "0100", "", "Mrs Live", self.MUM])
        self.assertEqual(self.gaps(first), [(2, "Chike Obi", "A phone number only.")])
        guardian = User.objects.get(guardian_account__contacts__value="+2348031230000")
        from django.utils import timezone

        from accounts.models import ContactChannel, GuardianContact
        from accounts.services import activate_guardian_links

        activate_guardian_links(guardian, self.stmarys)
        GuardianContact.objects.create(
            guardian=guardian.guardian_account, channel_type=ContactChannel.EMAIL, value="live@example.test",
            created_by=self.admin.user, verified_at=timezone.now(),
        )

        second = self.import_(["Dayo Obi", "JSS 1B", "0101", "", "Mrs Live", self.MUM])

        self.assertEqual(second["no_email"], [], "their sibling's mother is reachable by email already")

    def test_an_email_held_at_the_other_school_is_not_this_schools_to_know(self):
        """Mrs Live is live at Grace with an email. At St Mary's she has answered nobody, so a
        St Mary's import of her number still lists the child: the report would otherwise say
        which numbers belong to a parent with an email elsewhere."""
        self.live_guardian_with_an_email(
            self.grace, self.grace_child, "08031230000", "elsewhere@example.test"
        )

        body = self.import_(["Chike Obi", "JSS 1A", "0100", "", "Mrs Live", "08031230000"])

        self.assertEqual(self.gaps(body), [(2, "Chike Obi", "A phone number only.")])
        self.assertNotIn("elsewhere", str(body))

    def test_the_other_schools_import_is_judged_on_its_own_and_never_lists_ours(self):
        self.import_(["Chike Obi", "JSS 1A", "0100", "", "Mrs Obi", self.MUM])
        self.client.force_login(self.their_admin.user)
        from academics.models import ClassGroup

        with connected_to(self.grace):
            theirs = ClassGroup.objects.get().name
        response = self.client.post(
            FILE,
            {"file": upload("roll.xlsx", workbook(["Femi Eze", theirs, "G100", "", "Mrs Eze", "eze@example.test"]))},
            HTTP_HOST=THEIR_HOST,
        )

        self.assertEqual(response.json()["admitted"], 1)
        self.assertEqual(response.json()["no_email"], [], "and Chike Obi is not in Grace's report")
